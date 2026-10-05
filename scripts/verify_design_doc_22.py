from __future__ import annotations

import os
import zipfile
from pathlib import Path

from docx import Document


path = Path(os.environ["DESIGN_DOC_OUTPUT"])
doc = Document(path)
text = "\n".join(paragraph.text for paragraph in doc.paragraphs)

expected = {
    "version": "\u0434\u0438\u0437\u0430\u0439\u043d-\u0434\u043e\u043a\u0443\u043c\u0435\u043d\u0442 2.2",
    "acquisition": "\u041e\u0442\u0434\u0435\u043b\u044c\u043d\u043e\u0439 \u043c\u0438\u043d\u0438-\u0438\u0433\u0440\u044b \u043b\u043e\u0432\u043b\u0438 \u043d\u0435\u0442",
    "rank_vi": "VI \u0440\u0430\u043d\u0433 \u2014 \u0411\u043e\u0436\u0435\u0441\u0442\u0432\u0435\u043d\u043d\u0430\u044f \u0438\u0442\u0435\u0440\u0430\u0446\u0438\u044f",
    "armored_line": "\u041f\u0430\u043d\u0446\u0438\u0440\u043d\u044b\u0439 \u0440\u043e\u0433\u0430\u0447 I \u2014 Armored Stag Beetle I",
    "flail_line": "\u0416\u0443\u043a-\u043c\u043e\u043b\u043e\u0442\u0438\u043b\u044c\u0449\u0438\u043a II \u2014 Flail Beetle II",
    "resonance_v": "\u0420\u0435\u0437\u043e\u043d\u0430\u0442\u043e\u0440\u043d\u044b\u0439 \u0440\u043e\u0433\u0430\u0447 V \u2014 Resonance Stag Beetle V",
    "corrosion_v": "Corrosion Flail Beetle V",
    "region": "\u0422\u0435\u0440\u0440\u0430\u0441\u044b.",
    "states": "\u0418\u043d\u0436\u0435\u043d\u0435\u0440\u043d\u0430\u044f \u0430\u043d\u0442\u0438\u044d\u0445\u043e-\u0434\u0435\u0440\u0436\u0430\u0432\u0430",
    "resources": "\u0421\u0438\u0441\u0442\u0435\u043c\u0430 \u043d\u0430\u0433\u0440\u0430\u0434 \u0438\u0441\u043f\u043e\u043b\u044c\u0437\u0443\u0435\u0442 \u0447\u0435\u0442\u044b\u0440\u0435 \u043f\u0435\u0440\u0435\u0441\u0435\u043a\u0430\u044e\u0449\u0438\u0445\u0441\u044f \u0441\u043b\u043e\u044f \u0440\u0435\u0441\u0443\u0440\u0441\u043e\u0432",
}

for label, needle in expected.items():
    assert needle in text, f"Missing expected section: {label}"
    print(f"PASS {label}")

forbidden = {
    "old_region": "\u041a\u0430\u043c\u0435\u043d\u043d\u044b\u0435 \u0422\u0435\u0440\u0440\u0430\u0441\u044b",
    "old_wave": "\u041a\u0440\u0443\u0433\u043e\u0432\u0430\u044f \u0432\u043e\u043b\u043d\u0430",
    "old_stomper": "\u0420\u043e\u0433\u0430\u0447-\u0442\u043e\u043f\u0442\u0443\u043d",
}

for label, needle in forbidden.items():
    assert needle not in text, f"Forbidden stale text remains: {label}"
    print(f"PASS no_{label}")

with zipfile.ZipFile(path) as package:
    media = [name for name in package.namelist() if name.startswith("word/media/")]
assert len(media) == 16, f"Expected 16 embedded images, found {len(media)}"
print(f"PASS embedded_images={len(media)}")
print(f"PASS paragraphs={len(doc.paragraphs)}")
print(f"PASS bytes={path.stat().st_size}")
