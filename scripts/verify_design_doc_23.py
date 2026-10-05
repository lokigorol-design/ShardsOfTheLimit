from __future__ import annotations

import os
import zipfile
from pathlib import Path

from docx import Document


path = Path(os.environ["DESIGN_DOC_OUTPUT"])
doc = Document(path)
paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs]
text = "\n".join(paragraphs)

expected = {
    "version": "дизайн-документ 2.3",
    "echo_law": "Эхо — живые отголоски законов расколотого мира",
    "archon_rarity": "примерно один человек на тысячу",
    "amplifier_frame": "в оправу его резонансного кристалла",
    "three_active": "три стартовых Эхо",
    "recall_channel": "состояние канала отзыва",
    "shared_library": "80–90% способностей являются общими механическими архетипами",
    "two_talents": "максимум два таланта",
    "hex_grid": "гексагональную сетку",
    "battle_length": "5–10 минут",
    "divine": "одну масштабную способность на экспедицию",
    "no_theo": "Тео, Оро и Фау не возвращаются в формулы",
}

for label, needle in expected.items():
    assert needle in text, f"Missing expected update: {label}"
    print(f"PASS {label}")

forbidden = {
    "archon_unit": "Архонт физически присутствует на тактическом поле",
    "old_open_active_count": "Необходимо определить точное число одновременно активных Эхо",
    "old_mvp_archon": "таймлайн, Архонт, призыв и отзыв",
}

for label, needle in forbidden.items():
    assert needle not in text, f"Stale statement remains: {label}"
    print(f"PASS no_{label}")

assert paragraphs.index("2.35.1. Поле и длительность боя") < paragraphs.index("2.36. Визуальная эволюция Эхо")
assert paragraphs.index("2.38. Архонт, призыв и отзыв Эхо") < paragraphs.index("2.39. Террасный рогач: актуальная схема развития")
print("PASS section_order")

with zipfile.ZipFile(path) as package:
    media = [name for name in package.namelist() if name.startswith("word/media/")]
assert len(media) == 16, f"Expected 16 embedded images, found {len(media)}"

print(f"PASS embedded_images={len(media)}")
print(f"PASS paragraphs={len(doc.paragraphs)}")
print(f"PASS bytes={path.stat().st_size}")
